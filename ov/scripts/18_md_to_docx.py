#!/usr/bin/env python3
"""Convert the manuscript from Markdown to .docx.

Written rather than pandoc'd because pandoc is not installed here. Handles the
subset of Markdown the manuscript actually uses: ATX headings, pipe tables,
blockquotes, ordered/unordered lists, horizontal rules, and inline **bold**,
*italic* and `code`. Figures are embedded with numbered captions in place of the
figure-index table.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.environ.get('GIAMIA_ROOT',
                      os.path.abspath(os.path.join(os.path.dirname(__file__),
                                                   '..', '..')))
SRC = f'{ROOT}/ov/paper/manuscript.md'
FIGDIR = f'{ROOT}/ov/figures'
OUT = f'{ROOT}/ov/paper/manuscript.docx'

FIGS = [
    ('fig1_prompt_groups.png',
     'Prevalence-corrected Cohen’s κ (a) and best F1 at IoU 0.50 (b) by prompt '
     'group for six detectors, with the label-noise ceiling band and the '
     'expert–expert agreement line. Group D is not size-matched here; see Fig. 5.'),
    ('fig2_contrasts.png',
     'Paired prompt contrasts with 95 % bootstrap confidence intervals, six '
     'detectors. Every contrast changes sign between architecture families.'),
    ('fig3_annotation_protocol.png',
     'Box-area distributions of the two annotated classes. `weed` is annotated at '
     'instance level (median 4.0 % of frame), `sugarcane` at region level '
     '(median 30.7 %), a 7.8-fold difference (Instrument 1).'),
    ('fig4_false_alarms.png',
     'Fraction of the 869 expert-confirmed weed-free images producing at least one '
     'detection, as a function of confidence threshold. Solid: weed prompts '
     '(false alarms). Dashed: crop prompt (true positives). Markers show each '
     'run’s F1-optimal operating point, fixed on the annotated set beforehand.'),
    ('fig5_confound_reversal.png',
     'The crop–weed contrast (D − A) before and after matching ground-truth box '
     'size, using the same metric on the same 95 images. The sign reverses in '
     'three of four detectors. Proposed graphical abstract.'),
    ('fig6_sam3_agent.png',
     'SAM 3 prompted directly versus SAM 3 Agent, the MLLM wrapper its authors '
     'propose for queries beyond simple noun phrases, on both evaluation sets.'),
]

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)')


def add_runs(par, text):
    """Render inline bold / italic / code into runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            par.add_run(part)


def hrule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', 'BBBBBB')):
        bot.set(qn(k), v)
    b.append(bot); pPr.append(b)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{2,}:?', c or '-') for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        wr = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ''
            par = wr[j].paragraphs[0]
            add_runs(par, txt)
            for run in par.runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def main():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(10.5)

    lines = open(SRC, encoding='utf-8').read().split('\n')
    i, fig_i = 0, 0
    in_figsec = False

    while i < len(lines):
        ln = lines[i]

        # the figure-index table is replaced by the embedded figures
        if ln.startswith('## Figures'):
            in_figsec = True
            doc.add_page_break()
            doc.add_heading('Figures', level=1)
            for fn, cap in FIGS:
                fig_i += 1
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                path = os.path.join(FIGDIR, fn)
                if os.path.exists(path):
                    p.add_run().add_picture(path, width=Inches(6.0))
                cp = doc.add_paragraph()
                cr = cp.add_run(f'Fig. {fig_i}. ')
                cr.bold = True; cr.font.size = Pt(9)
                r2 = cp.add_run(cap); r2.font.size = Pt(9)
                cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                doc.add_paragraph()
            i += 1
            continue
        if in_figsec:
            if ln.startswith('## '):
                in_figsec = False
            else:
                i += 1
                continue

        if not ln.strip():
            i += 1; continue
        if ln.strip() == '---':
            hrule(doc); i += 1; continue

        m = re.match(r'^(#{1,4})\s+(.*)', ln)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading('', level=min(lvl, 4))
            add_runs(h, m.group(2))
            i += 1; continue

        if ln.lstrip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                rows.append(lines[i]); i += 1
            add_table(doc, rows); continue

        if ln.startswith('> '):
            block = []
            while i < len(lines) and lines[i].startswith('>'):
                block.append(lines[i].lstrip('>').strip()); i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.4)
            add_runs(p, ' '.join(x for x in block if x))
            for r in p.runs:
                r.italic = True
            continue

        m = re.match(r'^(\s*)([-*])\s+(.*)', ln)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, m.group(3)); i += 1; continue
        m = re.match(r'^(\s*)(\d+)\.\s+(.*)', ln)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(3)); i += 1; continue

        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(
                r'^(#{1,4}\s|\||>|\s*[-*]\s|\s*\d+\.\s)', lines[i]) and lines[i].strip() != '---':
            buf.append(lines[i].strip()); i += 1
        if buf:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, ' '.join(buf))

    doc.save(OUT)
    print(f'wrote {OUT} ({os.path.getsize(OUT)/1024:.0f} KB), {fig_i} figures embedded')


if __name__ == '__main__':
    main()
